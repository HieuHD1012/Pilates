"""File mẫu gửi studio phải khớp thứ script nhập liệu thật sự đọc.

Một file mẫu lệch tên cột còn tệ hơn không có file mẫu: studio điền xong 500
dòng, gửi lại, và mọi dòng đều báo "thiếu mã học viên". Chỗ lệch đó không lộ
ra cho tới ngày nhập thật.

Test này đọc tên cột **thẳng từ mã nguồn** của script thay vì chép lại danh
sách — một danh sách chép tay là bản sao thứ hai, và bản sao thứ hai sẽ là bản
trôi lệch.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scripts.import_initial_data import _rows

TEMPLATE = Path(__file__).resolve().parents[2] / "docs/templates/mau-nhap-du-lieu-ban-dau.xlsx"
IMPORTER = Path(__file__).resolve().parents[1] / "scripts/import_initial_data.py"

#: Sheet dữ liệu và hàm đọc nó.
SHEET_READERS = {
    "hoc_vien": "_import_students",
    "hlv": "_import_trainers",
    "goi_tap": "_import_packages",
    "lich_lop": "_import_sessions",
}


def _columns_the_importer_reads(function_name: str) -> set[str]:
    body = IMPORTER.read_text().split(f"def {function_name}(")[1].split("\ndef ")[0]
    return set(re.findall(r'row\.get\("([a-z_]+)"\)', body))


@pytest.mark.parametrize(("sheet", "reader"), SHEET_READERS.items())
def test_template_offers_every_column_the_importer_reads(sheet: str, reader: str) -> None:
    workbook = load_workbook(TEMPLATE, read_only=True)
    assert sheet in workbook.sheetnames

    headers = {cell.value for cell in next(workbook[sheet].iter_rows(max_row=1))}
    wanted = _columns_the_importer_reads(reader)

    assert wanted - headers == set(), f"File mẫu thiếu cột: {sorted(wanted - headers)}"
    assert headers - wanted == set(), f"File mẫu thừa cột không ai đọc: {sorted(headers - wanted)}"


def test_template_ships_with_no_example_rows() -> None:
    """Các sheet dữ liệu cố ý chỉ có dòng tiêu đề.

    Một dòng ví dụ quên xoá sẽ thành một học viên có thật trong hệ thống, kèm
    một gói tập và một số dư mở sổ. Hướng dẫn nằm ở sheet riêng.
    """
    workbook = load_workbook(TEMPLATE, data_only=True, read_only=True)
    for sheet in SHEET_READERS:
        assert _rows(workbook, sheet) == [], f"Sheet {sheet} có dòng dữ liệu sẵn."
    assert "huong_dan" in workbook.sheetnames
