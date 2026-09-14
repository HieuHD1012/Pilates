"""Nhập dữ liệu ban đầu chạy lại nhiều lần không nhân đôi số dư.

Kịch bản đã lường trước, và nó không phải giả thiết: 12/11 trên PROD, dòng
340/500 lỗi vì một số điện thoại thừa khoảng trắng đụng `UNIQUE`; tiến trình
chết; người vận hành sửa file rồi chạy lại cả file. Nếu lần chạy thứ hai ghi
tiếp 339 dòng đầu thì **mỗi gói có thêm một bút toán mở sổ và số dư nhân đôi**.
Sổ buổi append-only nên gỡ ra phải ghi 339 bút toán đối ứng bằng tay — trong
tuần nghiệm thu.

Điều được kiểm ở đây không phải "script chạy xong" mà là **chạy hai lần cho
đúng cùng một trạng thái**.
"""

from __future__ import annotations

import hashlib
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.domain.rules import TIMEZONE, LedgerReason, Role, UserStatus
from app.models.money import CreditLedger, StudentPackage
from app.models.people import Student, Trainer
from app.models.scheduling import ClassSession, ImportRun
from app.models.user import User
from app.services.ledger_invariants import assert_ledger_is_sound
from scripts.import_initial_data import main, run_import
from tests.conftest import make_user

SOURCE = "studio-thang-11"


def _package_row(
    external_ref: str, student_ref: str, name: str, class_type: str, credits: int
) -> list:
    return [
        external_ref,
        student_ref,
        name,
        class_type,
        credits,
        date(2026, 9, 1),
        date(2026, 12, 1),
        5000000,
    ]


def _workbook(tmp_path: Path, *, student_rows=None, package_rows=None) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "hoc_vien"
    sheet.append(["ma_hoc_vien", "ho_ten", "so_dien_thoai", "email", "ghi_chu"])
    for row in student_rows or [
        ["HV001", "Nguyễn Thị Lan", "0900 000 055", "lan@example.com", "Khách quen"],
        ["HV002", "Trần Văn Minh", "0900000056", None, None],
    ]:
        sheet.append(row)

    trainers = workbook.create_sheet("hlv")
    trainers.append(["ho_ten", "so_dien_thoai", "email"])
    trainers.append(["HLV Ngọc Mai", "0900000077", "mai@example.com"])

    packages = workbook.create_sheet("goi_tap")
    packages.append(
        [
            "ma_goi",
            "ma_hoc_vien",
            "ten_goi",
            "loai_lop",
            "so_buoi_con_lai",
            "ngay_bat_dau",
            "ngay_het_han",
            "gia",
        ]
    )
    for row in package_rows or [
        _package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 13),
        _package_row("GT002", "HV002", "Gói Private 10", "PRIVATE", 4),
    ]:
        packages.append(row)

    sessions = workbook.create_sheet("lich_lop")
    sessions.append(["ma_lop", "ten_hlv", "bat_dau", "ket_thuc", "loai_lop", "suc_chua"])
    sessions.append(
        [
            "LOP001",
            "HLV Ngọc Mai",
            datetime(2026, 12, 1, 6, 0),
            datetime(2026, 12, 1, 7, 0),
            "GROUP",
            6,
        ]
    )

    path = tmp_path / "du-lieu.xlsx"
    workbook.save(path)
    return path


@pytest.fixture
def admin(db: Session) -> User:
    return make_user(db, Role.ADMIN)


def _counts(db: Session) -> dict[str, int]:
    return {
        "students": db.scalar(select(func.count()).select_from(Student)),
        "packages": db.scalar(select(func.count()).select_from(StudentPackage)),
        "sessions": db.scalar(select(func.count()).select_from(ClassSession)),
        "trainers": db.scalar(select(func.count()).select_from(Trainer)),
        "ledger": db.scalar(select(func.count()).select_from(CreditLedger)),
    }


def _balances(db: Session) -> dict[str, int]:
    rows = db.execute(
        select(
            StudentPackage.external_ref,
            func.coalesce(func.sum(CreditLedger.delta), 0),
        )
        .outerjoin(CreditLedger, CreditLedger.student_package_id == StudentPackage.id)
        .group_by(StudentPackage.external_ref)
    ).all()
    return {ref: int(total) for ref, total in rows}


def test_running_the_same_file_twice_changes_nothing(
    db: Session, admin, tmp_path: Path
) -> None:
    path = _workbook(tmp_path)

    first, failures = run_import(db, path=path, source=SOURCE, actor_user_id=admin.id)
    db.commit()
    assert failures == []
    assert (first.students_created, first.packages_created) == (2, 2)

    after_first = _counts(db)
    balances_first = _balances(db)

    second, failures = run_import(db, path=path, source=SOURCE, actor_user_id=admin.id)
    db.commit()

    assert failures == []
    assert second.students_created == 0
    assert second.packages_created == 0
    assert second.sessions_created == 0
    assert (second.students_skipped, second.packages_skipped) == (2, 2)
    assert _counts(db) == after_first
    assert _balances(db) == balances_first == {"GT001": 13, "GT002": 4}
    assert_ledger_is_sound(db)


def test_a_resumed_run_after_a_mid_file_failure_does_not_double_anything(
    db: Session, admin, tmp_path: Path
) -> None:
    """Lần chạy đầu chết giữa chừng, người vận hành sửa file rồi chạy lại cả file.

    Đây là đúng đường đi tới sự cố: phần đã vào không được vào lần nữa, phần
    còn thiếu phải vào đủ.
    """
    partial = _workbook(
        tmp_path,
        student_rows=[
            ["HV001", "Nguyễn Thị Lan", "0900 000 055", "lan@example.com", None]
        ],
        package_rows=[_package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 13)],
    )
    run_import(db, path=partial, source=SOURCE, actor_user_id=admin.id)
    db.commit()

    second_dir = tmp_path / "lan-hai"
    second_dir.mkdir()
    full = _workbook(second_dir)
    report, failures = run_import(db, path=full, source=SOURCE, actor_user_id=admin.id)
    db.commit()

    assert failures == []
    assert report.students_created == 1
    assert report.students_skipped == 1
    assert report.packages_created == 1
    assert report.packages_skipped == 1
    assert _balances(db) == {"GT001": 13, "GT002": 4}

    opening_entries = db.scalar(
        select(func.count())
        .select_from(CreditLedger)
        .where(CreditLedger.reason_code == LedgerReason.PACKAGE_SOLD)
    )
    assert opening_entries == 2
    assert_ledger_is_sound(db)


def test_imported_credits_enter_the_ledger_not_the_cached_balance(
    db: Session, admin, tmp_path: Path
) -> None:
    """Gói cũ vào bằng **bút toán mở sổ**, không phải một lần gán số dư.

    Ghi thẳng `balance_cached` là cách hai biểu diễn số dư lệch nhau ngay từ
    ngày đầu chạy thật — và phép đối soát sẽ báo xanh, vì nó chỉ so số dư với
    chính nó.
    """
    run_import(db, path=_workbook(tmp_path), source=SOURCE, actor_user_id=admin.id)
    db.commit()

    package = db.scalar(
        select(StudentPackage).where(StudentPackage.external_ref == "GT001")
    )
    entries = list(
        db.scalars(
            select(CreditLedger).where(CreditLedger.student_package_id == package.id)
        )
    )
    assert [entry.reason_code for entry in entries] == [LedgerReason.PACKAGE_SOLD]
    assert entries[0].delta == 13
    assert package.balance_cached == 13
    # File không khai số buổi gốc, nên bản ghi nhận đúng thứ biết được: 13.
    # Bịa ra 20 là nói rằng học viên đã tập 7 buổi — điều không ai kiểm được.
    assert package.credits_snapshot == 13


def test_a_balance_that_disagrees_with_the_file_aborts_the_import(
    db: Session, admin, tmp_path: Path
) -> None:
    """Đối chiếu phải so với **con số trong file**, không so sổ với chính nó.

    Dựng lệch bằng cách cho hai dòng cùng mã gói khai hai số buổi khác nhau:
    dòng đầu tạo gói với 13 buổi, dòng sau thấy gói đã có nên bỏ qua nhưng vẫn
    đòi 99 buổi. Sổ vẫn cân — và vẫn phải bị chặn.
    """
    path = _workbook(
        tmp_path,
        package_rows=[
            _package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 13),
            _package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 99),
        ],
    )
    _, failures = run_import(db, path=path, source=SOURCE, actor_user_id=admin.id)
    db.rollback()

    assert any("99" in message for message in failures), failures


def test_imported_accounts_have_no_password(db: Session, admin, tmp_path: Path) -> None:
    """Tài khoản nhập từ Excel **không có mật khẩu** và chờ kích hoạt.

    Không bao giờ sinh mật khẩu suy ra được từ số điện thoại: số điện thoại lộ
    qua form tư vấn công khai và danh bạ studio.
    """
    run_import(db, path=_workbook(tmp_path), source=SOURCE, actor_user_id=admin.id)
    db.commit()

    imported = list(
        db.scalars(
            select(User).where(
                User.email.in_(["lan@example.com", "mai@example.com"])
            )
        )
    )
    assert len(imported) == 2
    assert all(user.password_hash is None for user in imported)
    assert all(user.status is UserStatus.PENDING_ACTIVATION for user in imported)


def test_a_student_without_an_email_gets_no_account(
    db: Session, admin, tmp_path: Path
) -> None:
    """Không email thì không tài khoản — chứ không bịa một địa chỉ suy ra được."""
    run_import(db, path=_workbook(tmp_path), source=SOURCE, actor_user_id=admin.id)
    db.commit()
    student = db.scalar(select(Student).where(Student.external_ref == "HV002"))
    assert student.user_id is None


def test_phone_numbers_are_normalised_before_they_hit_the_unique_index(
    db: Session, admin, tmp_path: Path
) -> None:
    """`0900 000 055` và `0900000055` là cùng một người.

    Chính khoảng trắng thừa này là thứ làm lần chạy đầu chết giữa chừng ở kịch
    bản đã lường trước.
    """
    run_import(db, path=_workbook(tmp_path), source=SOURCE, actor_user_id=admin.id)
    db.commit()
    student = db.scalar(select(Student).where(Student.external_ref == "HV001"))
    assert student.phone == "0900000055"


def test_class_times_are_read_as_studio_time(db: Session, admin, tmp_path: Path) -> None:
    """Giờ trong file là giờ Nha Trang, không phải UTC.

    Đọc nhầm là dời cả lịch lớp đi 7 tiếng, và không ràng buộc nào phản đối vì
    07:00 vẫn là một giờ hợp lệ.
    """
    run_import(db, path=_workbook(tmp_path), source=SOURCE, actor_user_id=admin.id)
    db.commit()

    class_session = db.scalar(
        select(ClassSession).where(ClassSession.external_ref == "LOP001")
    )
    local = class_session.starts_at.astimezone(TIMEZONE)
    assert (local.hour, local.minute) == (6, 0)
    assert local.date() == date(2026, 12, 1)
    assert class_session.ends_at - class_session.starts_at == timedelta(hours=1)


# --- Một dòng hỏng không được giết cả file -----------------------------------


def test_two_rows_sharing_a_phone_report_the_row_and_import_the_rest(
    db: Session, admin, tmp_path: Path
) -> None:
    """Mẹ và con dùng chung số điện thoại — chuyện rất thường ở studio Việt Nam.

    `student.phone` là `UNIQUE`, nên dòng thứ hai không vào được. Điều phải xảy
    ra là **một dòng `problems` chỉ đúng dòng nào**, không phải một traceback
    Python giữa buổi nhập liệu, để lại người vận hành không biết dòng nào hỏng
    và không dòng nhật ký nào.
    """
    path = _workbook(
        tmp_path,
        student_rows=[
            ["HV001", "Nguyễn Thị Lan", "0900 111 222", None, None],
            ["HV002", "Nguyễn Minh Anh", "0900111222", None, None],
            ["HV003", "Trần Văn Minh", "0900000099", None, None],
        ],
        package_rows=[_package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 13)],
    )
    report, failures = run_import(db, path=path, source=SOURCE, actor_user_id=admin.id)
    db.commit()

    assert report.students_created == 2
    assert len(report.problems) == 1
    assert "hoc_vien!3" in report.problems[0]
    assert "số điện thoại" in report.problems[0]
    assert failures == []

    names = set(db.scalars(select(Student.full_name)))
    assert names == {"Nguyễn Thị Lan", "Trần Văn Minh"}
    assert _balances(db) == {"GT001": 13}
    assert_ledger_is_sound(db)


def test_two_rows_sharing_an_email_still_import_both_people(
    db: Session, admin, tmp_path: Path
) -> None:
    """Người thứ hai vẫn được nhập, chỉ là **chưa có tài khoản**.

    `student.user_id` là `UNIQUE` nên chỉ một người giữ được tài khoản. Bỏ luôn
    cả dòng vì lý do đó là mất một học viên có thật khỏi hệ thống.
    """
    path = _workbook(
        tmp_path,
        student_rows=[
            ["HV001", "Nguyễn Thị Lan", "0900111222", "chung@example.com", None],
            ["HV002", "Nguyễn Minh Anh", "0900111333", "chung@example.com", None],
        ],
        package_rows=[_package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 13)],
    )
    report, failures = run_import(db, path=path, source=SOURCE, actor_user_id=admin.id)
    db.commit()

    assert report.students_created == 2
    assert failures == []
    assert any("chưa có tài khoản" in message for message in report.problems)

    with_account = db.scalar(
        select(func.count()).select_from(Student).where(Student.user_id.is_not(None))
    )
    assert with_account == 1


# --- `main()`: chạy thử, ghi thật, và dấu vết --------------------------------


def _run_main(path: Path, *extra: str) -> int:
    return main(["--file", str(path), "--source", SOURCE, *extra])


def test_a_rehearsal_leaves_a_trace_but_no_data(
    db: Session, admin, tmp_path: Path
) -> None:
    """Chạy thử **không** ghi dữ liệu nhưng **có** ghi nhật ký.

    Nhật ký chạy không được chung số phận với dữ liệu nó mô tả: ghi chung
    transaction thì lần chạy thử rollback luôn dòng nhật ký, và yêu cầu "diễn
    tập đầy đủ trước ngày nhập thật" không để lại gì để chứng minh đã diễn tập.
    """
    path = _workbook(tmp_path)
    assert _run_main(path) == 0

    db.expire_all()
    assert _counts(db)["students"] == 0
    assert _counts(db)["packages"] == 0

    runs = list(db.scalars(select(ImportRun)))
    assert len(runs) == 1
    assert runs[0].dry_run is True
    assert runs[0].source == SOURCE
    assert runs[0].file_hash == hashlib.sha256(path.read_bytes()).hexdigest()
    assert '"created": 2' in runs[0].summary


def test_commit_writes_the_data_and_a_second_trace(
    db: Session, admin, tmp_path: Path
) -> None:
    path = _workbook(tmp_path)
    assert _run_main(path) == 0
    assert _run_main(path, "--commit") == 0

    db.expire_all()
    assert _counts(db)["students"] == 2
    assert _balances(db) == {"GT001": 13, "GT002": 4}

    runs = list(db.scalars(select(ImportRun).order_by(ImportRun.id)))
    assert [item.dry_run for item in runs] == [True, False]
    assert len({item.file_hash for item in runs}) == 1
    assert_ledger_is_sound(db)


def test_a_balance_mismatch_rolls_back_data_but_keeps_the_trace(
    db: Session, admin, tmp_path: Path
) -> None:
    """Lần nhập hỏng vẫn phải để lại bằng chứng nó đã chạy và hỏng vì gì."""
    path = _workbook(
        tmp_path,
        package_rows=[
            _package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 13),
            _package_row("GT001", "HV001", "Gói 20 buổi", "GROUP", 99),
        ],
    )
    assert _run_main(path, "--commit") == 2

    db.expire_all()
    assert _counts(db)["students"] == 0
    runs = list(db.scalars(select(ImportRun)))
    assert len(runs) == 1
    assert "99" in runs[0].summary
